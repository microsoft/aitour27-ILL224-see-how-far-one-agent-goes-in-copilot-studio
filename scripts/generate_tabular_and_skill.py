#!/usr/bin/env python3
"""Generate tabular, skill, learner VM, and manifest artifacts for ILL224."""

from __future__ import annotations

import csv
import hashlib
import shutil
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parent
PACKAGE = ROOT / "package"
SEED = PACKAGE / "SharePoint-Lists" / "seed-data"
VM = PACKAGE / "LearnerVM"
OPTIONAL = PACKAGE / "SharePoint" / "optional"

NAVY = "16324F"
TEAL = "007F73"
MINT = "DDF3EC"
GOLD = "F2B134"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="C5D0D8")


INVOICES = [
    {
        "invoice_id": "ARB-260814",
        "supplier_id": "SUP-1042",
        "supplier_name": "Astor Ridge Biologics",
        "purchase_order": "PO-88431",
        "contract_id": "CMO-2026-041",
        "invoice_date": "2026-08-14",
        "batch_id": "CDV-AX47",
        "currency": "USD",
        "invoice_total": 217700,
        "lines": [
            (1, "Commercial manufacturing batch", 1, 185000, 185000, "Section 4.1"),
            (2, "Expedited production premium", 1, 14800, 14800, "Section 4.2"),
            (3, "Cold-chain freight", 1, 12200, 12200, "Section 4.3"),
            (4, "Validation documentation package", 1, 5700, 5700, ""),
        ],
    },
    {
        "invoice_id": "NWT-260821",
        "supplier_id": "SUP-1088",
        "supplier_name": "Northwind Therapeutics",
        "purchase_order": "PO-89214",
        "contract_id": "PKG-2026-052",
        "invoice_date": "2026-08-21",
        "batch_id": "CDV-LB22",
        "currency": "USD",
        "invoice_total": 122300,
        "lines": [
            (1, "Packaging and labeling production run", 1, 96000, 96000, "Section 4.1"),
            (2, "Serialized label setup", 1, 12000, 12000, "Section 4.2"),
            (3, "Standard freight", 1, 4800, 4800, "Section 4.3"),
            (4, "Rush changeover fee", 1, 9500, 9500, ""),
        ],
    },
    {
        "invoice_id": "CBP-260825",
        "supplier_id": "SUP-1114",
        "supplier_name": "Contoso BioProcessing",
        "purchase_order": "PO-89501",
        "contract_id": "CLN-2026-063",
        "invoice_date": "2026-08-25",
        "batch_id": "CDV-CL19",
        "currency": "USD",
        "invoice_total": 89500,
        "lines": [
            (1, "Sterile vial preparation and filling", 1, 65000, 65000, "Section 3.1"),
            (2, "Release testing coordination", 1, 18000, 18000, "Section 3.2"),
            (3, "Environmental monitoring", 1, 4500, 4500, "Section 3.4"),
            (4, "Deviation handling", 1, 2000, 2000, ""),
        ],
    },
    {
        "invoice_id": "FCC-260829",
        "supplier_id": "SUP-1150",
        "supplier_name": "Fabrikam Cold Chain",
        "purchase_order": "PO-89712",
        "contract_id": "LOG-2026-074",
        "invoice_date": "2026-08-29",
        "batch_id": "CDV-FC08",
        "currency": "USD",
        "invoice_total": 67200,
        "lines": [
            (1, "Refrigerated transport and storage", 1, 51000, 51000, "Section 5.1"),
            (2, "Cold-chain monitoring and reporting", 1, 8200, 8200, "Section 5.3"),
            (3, "Fuel surcharge", 1, 5000, 5000, "Section 5.4"),
            (4, "Delivery validation package", 1, 3000, 3000, ""),
        ],
    },
]


def ensure_directories() -> None:
    legacy_workflow = PACKAGE / "Workflow"
    if legacy_workflow.exists():
        shutil.rmtree(legacy_workflow)
    for directory in [SEED, VM / "case-inputs", VM / "skills", VM / "evaluation", OPTIONAL]:
        directory.mkdir(parents=True, exist_ok=True)


def style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=GRID)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_view.showGridLines = False


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def generate_invoice_workbook() -> None:
    invoice = INVOICES[0]
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Invoice"
    summary.append(["CALDOVA SUPPLIER INVOICE - FICTIONAL LAB DATA", ""])
    summary.merge_cells("A1:B1")
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].font = Font(color=WHITE, bold=True, size=13)
    summary["A1"].alignment = Alignment(horizontal="center")
    fields = [
        ("Invoice ID", invoice["invoice_id"]), ("Supplier ID", invoice["supplier_id"]),
        ("Supplier", invoice["supplier_name"]), ("Purchase Order", invoice["purchase_order"]),
        ("Contract ID", invoice["contract_id"]), ("Invoice Date", invoice["invoice_date"]),
        ("Batch ID", invoice["batch_id"]), ("Currency", invoice["currency"]),
        ("Invoice Total", invoice["invoice_total"]), ("Payment Status", "Pending review"),
    ]
    for label, value in fields:
        summary.append([label, value])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 46
    for cell in summary["A"][1:]:
        cell.font = Font(bold=True, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=MINT)
    summary["B10"].number_format = "$#,##0.00"
    summary.sheet_view.showGridLines = False
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 1

    lines = workbook.create_sheet("Line Items")
    lines.append(["Sequence", "Description", "Quantity", "Unit Price", "Amount", "Claimed Contract Reference"])
    for item in invoice["lines"]:
        lines.append(list(item))
    style_sheet(lines, [12, 38, 12, 16, 16, 28])
    for row in range(2, lines.max_row + 1):
        lines.cell(row, 4).number_format = "$#,##0.00"
        lines.cell(row, 5).number_format = "$#,##0.00"
    add_table(lines, "AstorInvoiceLines")
    lines.append(["", "Invoice total", "", "", f"=SUM(E2:E{lines.max_row})", ""])
    lines.cell(lines.max_row, 2).font = Font(bold=True)
    lines.cell(lines.max_row, 5).font = Font(bold=True, color=TEAL)
    lines.cell(lines.max_row, 5).number_format = "$#,##0.00"
    lines.page_setup.orientation = "landscape"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output = VM / "case-inputs" / "Caldova_Astor_Ridge_Invoice.xlsx"
    workbook.save(output)


def write_csv(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def generate_seed_data() -> None:
    write_csv(SEED / "Suppliers.csv",
              ["Title", "Supplier ID", "Status", "Risk Tier", "Contract ID", "Payment Terms", "Eligible for Standard Controls"],
              [["Astor Ridge Biologics", "SUP-1042", "Active", "Medium", "CMO-2026-041", "Net 60", "Yes"],
               ["Northwind Therapeutics", "SUP-1088", "Active", "Low", "PKG-2026-052", "Net 45", "Yes"],
               ["Contoso BioProcessing", "SUP-1114", "Conditional", "High", "CLN-2026-063", "Net 60", "No"],
               ["Fabrikam Cold Chain", "SUP-1150", "Active", "Medium", "LOG-2026-074", "Net 45", "Yes"]])
    write_csv(SEED / "SupplierInvoices.csv",
              ["Title", "Supplier ID", "Supplier Name", "Purchase Order", "Contract ID", "Invoice Date", "Batch or Lot ID", "Currency", "Invoice Total", "Status"],
              [[item["invoice_id"], item["supplier_id"], item["supplier_name"], item["purchase_order"], item["contract_id"], item["invoice_date"], item["batch_id"], item["currency"], item["invoice_total"], "Pending Review"] for item in INVOICES])
    invoice_lines = []
    for invoice in INVOICES:
        for sequence, description, quantity, unit_price, amount, reference in invoice["lines"]:
            invoice_lines.append([f"{invoice['invoice_id']}-{sequence}", invoice["invoice_id"], sequence, description, quantity, unit_price, amount, reference])
    write_csv(SEED / "InvoiceLines.csv",
              ["Title", "Invoice ID", "Sequence", "Description", "Quantity", "Unit Price", "Amount", "Claimed Contract Reference"],
              invoice_lines)
    write_csv(SEED / "QualityEvents.csv",
              ["Title", "Supplier ID", "Batch or Lot ID", "Event Status", "Active Quality Hold", "Product Impact", "Summary"],
              [["QE-2026-118", "SUP-1042", "CDV-AX47", "Open", "No", "No", "Temperature logger showed a short excursion during transit. Quality review found no product impact; corrective-action tracking remains open."],
               ["QE-2026-127", "SUP-1114", "CDV-CL19", "Open", "Yes", "Under Review", "Release-testing deviation; Caldova Quality explicitly issued an active hold."],
               ["QE-2026-132", "SUP-1088", "CDV-LB22", "Closed", "No", "No", "Packaging lot records and temperature logs were reviewed with no exception or product impact."],
               ["QE-2026-141", "SUP-1150", "CDV-FC08", "Closed", "No", "No", "Cold-chain delivery remained within the validated temperature range with no product impact."]])
    write_csv(SEED / "InvoiceReviewRequests.csv",
              ["Title", "Invoice ID", "Request Notes", "Status", "Agent Response", "Error Detail"], [])
    write_csv(SEED / "InvoiceReviewLog.csv",
              ["Title", "Invoice ID", "Supported Amount", "Disputed Amount", "Human Review Required", "Recommendation", "Review Source", "Review Document Name"], [])


def generate_optional_workbook() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Supplier Portfolio"
    sheet.append(["Supplier ID", "Supplier Name", "Status", "Approved Services", "Risk Tier", "Contract ID", "Payment Terms", "Active Quality Hold", "Quality Event"])
    rows = [
        ["SUP-1042", "Astor Ridge Biologics", "Active", "Commercial batch manufacturing and cold-chain shipping", "Medium", "CMO-2026-041", "Net 60", "No", "QE-2026-118"],
        ["SUP-1088", "Northwind Therapeutics", "Active", "Packaging, labeling, and serialized label setup", "Low", "PKG-2026-052", "Net 45", "No", "QE-2026-132"],
        ["SUP-1114", "Contoso BioProcessing", "Conditional", "Clinical batch manufacturing and release testing", "High", "CLN-2026-063", "Net 60", "Yes", "QE-2026-127"],
        ["SUP-1150", "Fabrikam Cold Chain", "Active", "Validated cold-chain logistics", "Medium", "LOG-2026-074", "Net 45", "No", "QE-2026-141"],
    ]
    for row in rows:
        sheet.append(row)
    style_sheet(sheet, [14, 27, 14, 50, 12, 18, 16, 20, 18])
    sheet.page_setup.orientation = "landscape"
    add_table(sheet, "SupplierPortfolio")
    workbook.save(OPTIONAL / "Caldova_Multi_Case_Supplier_Master.xlsx")


def generate_evaluation_template() -> None:
    comments = [
        "# Import conversations to test your agent.",
        "#",
        "# Limitations",
        "# - 8 question-and-answer pairs max per conversation.",
        "# - 100 conversations max.",
        "# - 500 characters max per question, including spaces.",
        "#",
        "# Imported columns",
        "# conversationNumber - Identifies each conversation. All questions and responses with the same conversation number will run as a single test case against the agent.",
        "# question - The user prompt that the agent will respond to.",
        "# response - The reference agent reply. This field is optional. The agent response isn't compared to this reference answer.",
        "#",
        "# Test methods",
        "# - Test methods are not included in this template. You can select them after importing the test cases.",
        "# - By default, the 'General quality' test method is added to the imported test set.",
        "#",
        "# For more details, refer to the documentation: https://go.microsoft.com/fwlink/?linkid=2335991",
        "#",
    ]
    rows = [
        ["1", "Review Astor Ridge invoice ARB-260814.", "Report $204,600 supported, $13,100 disputed, $0 reconciliation difference, and human review required."],
        ["1", "Is the full expedite premium supported?", "No. The contract permits $7,400; the additional $7,400 is unsupported."],
        ["1", "Does open event QE-2026-118 mean the batch is on hold?", "No. The event is open, but Active Quality Hold is No."],
        ["2", "Under Caldova policy, does a disputed amount of exactly $10,000 require human review based on the amount alone?", "No. Only a disputed amount greater than $10,000 triggers the amount threshold. Another policy condition could still require human review."],
    ]
    evaluation_path = VM / "evaluation" / "Supplier-Assurance-Baseline.csv"
    with evaluation_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        for comment in comments:
            writer.writerow([comment])
        writer.writerow(["conversationNumber", "question", "response"])
        writer.writerows(rows)
    (VM / "evaluation" / "README.md").write_text(
        "# Evaluation Release Gate\n\n"
        "`Supplier-Assurance-Baseline.csv` uses the Copilot Studio conversation import template. Conversation 1 tests a contextual Astor Ridge review and its grounding guardrails. Conversation 2 verifies the policy boundary for a disputed amount of exactly $10,000.\n\n"
        "The imported `response` values are reference replies for reviewers; Copilot Studio does not compare agent output directly with them. The import adds the General quality test method by default. Select any additional test methods in Copilot Studio after import.\n",
        encoding="utf-8",
    )


def generate_skill_package() -> None:
    skill_zip = VM / "skills" / "invoice-review-decision.zip"
    with zipfile.ZipFile(skill_zip, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(REPO / "supplier-policy-triage" / "SKILL.md", "SKILL.md")
        archive.write(REPO / "supplier-policy-triage" / "calculate_supplier_review.py", "calculate_supplier_review.py")


def copy_learner_materials() -> None:
    instructions_target = VM / "instructions"
    if instructions_target.exists():
        shutil.rmtree(instructions_target)
    shutil.copytree(ROOT / "lab" / "instructions", instructions_target)
    shutil.copy2(ROOT / "lab" / "README.md", VM / "README.md")


def write_manifest() -> None:
    package_files = sorted(
        path
        for path in PACKAGE.rglob("*")
        if path.is_file()
        and path.name != "checksums.sha256"
        and not any(part.startswith(".") for part in path.relative_to(PACKAGE).parts)
    )
    lines = []
    for file_path in package_files:
        digest = hashlib.sha256(file_path.read_bytes()).hexdigest()
        lines.append(f"{digest}  {file_path.relative_to(PACKAGE).as_posix()}")
    (PACKAGE / "checksums.sha256").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    ensure_directories()
    generate_invoice_workbook()
    generate_seed_data()
    generate_optional_workbook()
    generate_evaluation_template()
    generate_skill_package()
    copy_learner_materials()
    write_manifest()


if __name__ == "__main__":
    main()
